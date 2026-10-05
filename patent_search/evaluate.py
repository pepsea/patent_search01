"""プログラム3: 特許本文と「調べたいこと」の関連性を、ローカル LLM で評価する。

作業の流れ:
 1. keyword_snippets: 本文から、調べたいことの関連語の周辺を抜粋する(本文が長いため)
 2. build_prompt    : 調べたいこと + 判定基準 + 特許の本文(抜粋)から、LLM に渡す文章を作る
 3. バックエンド    : ollama_backend(この PC) / guidance_backend(サーバー)で LLM に判定させる
 4. check_evidence  : LLM が出した「根拠の引用」が原文に実在するかを確認する
 5. evaluate_table  : 上を全件に繰り返し、関連度の高い順の結果表を返す(発明内容概要・出願企業も付ける)
 6. save_evaluation_excel: 結果表を Excel に保存する(文献番号は Google Patents へのリンク)

- 判定結果は JSON スキーマで制約する（Ollama: format 指定 / Hugging Face: guidance の json 文法）。
  どちらのバックエンドでも同じプロンプト・同じスキーマなので、サーバー移植時はバックエンドだけ差し替える。
- LLM が出した「根拠の引用」が本当に原文にあるかを機械的に確認し、結果に付ける（幻覚の検出）。
"""

from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Callable

SCHEMA = {
    "type": "object",
    "properties": {
        "score": {"type": "integer", "enum": [0, 1, 2, 3]},
        "judgement": {"type": "string", "enum": ["直接関連", "関連あり", "わずかに関連", "無関係"]},
        "reason": {"type": "string", "maxLength": 300},
        "summary": {"type": "string", "maxLength": 220},
        "evidence": {"type": "array", "items": {"type": "string", "maxLength": 200}, "maxItems": 3},
    },
    "required": ["score", "judgement", "reason", "summary", "evidence"],
    "additionalProperties": False,
}

SYSTEM = "あなたは特許調査の専門家です。与えられた特許の本文が、指定された調査テーマに関連するかを、根拠を示して厳密に判定します。"

USER_TEMPLATE = """# 調査テーマ
{name}

## テーマの定義
{definition}

## 関連語（同義語・日英表記）
{keywords}

## 含める条件（関連とみなす）
{include}

## 除外する条件（関連とみなさない）
{exclude}

# 判定基準
- 3 (直接関連): 特許の中心的な内容（請求項・要約）が、このテーマの手法・対象そのものである
- 2 (関連あり): テーマの手法・対象を、明細書や実施例で実際に使用・言及している（請求項の中心ではない）
- 1 (わずかに関連): 関連語が出るが、一般的な背景説明や列挙の一部にとどまる
- 0 (無関係): 関連しない

# 対象特許
文献番号: {doc_no}
名称: {title}

## 要約
{abstract}

## 請求項
{claims}

## 明細書（関連語の周辺の抜粋 / 先頭部分）
{description}

# 出力
JSON のみを出力する。summary には、この特許の発明の内容（何を、どうする発明か）を、専門外の人にも分かる日本語で100〜150字にまとめる（要約と請求項に基づき、本文にないことは書かない）。
evidence には、上の本文からそのまま抜き出した短い引用（最大3件、各200字以内、改変禁止）を入れる。
根拠が本文にない場合は evidence を空配列にし、score は 0 か 1 にする。"""


@dataclass
class Topic:
    name: str
    definition: str
    keywords: list[str] = field(default_factory=list)
    include: str = ""
    exclude: str = ""


def nfkc(s: str) -> str:
    return unicodedata.normalize("NFKC", s)


def _keyword_patterns(keywords: list[str]) -> list[re.Pattern]:
    """通常の語は大文字小文字を区別しない。SHAPE・NAI・DMS のような大文字だけの短い略語は、
    shape(形状)・naive・DGSHAPE 等への誤一致を避けるため、大文字小文字を区別し前後に英数字がある場合は除外する。"""
    ci, strict = [], []
    for k in map(nfkc, keywords):
        if re.fullmatch(r"[A-Z0-9]{2,6}", k):
            strict.append(r"(?<![A-Za-z0-9])" + re.escape(k) + r"(?![A-Za-z0-9])")
        else:
            ci.append(re.escape(k))
    pats = []
    if ci:
        pats.append(re.compile("|".join(ci), re.IGNORECASE))
    if strict:
        pats.append(re.compile("|".join(strict)))
    return pats


# 作業: 関連語の前後を抜粋する。近い箇所は 1 つにまとめ、全ヒット数も返す。
def keyword_snippets(text: str, keywords: list[str], width: int = 150, max_snippets: int = 8) -> tuple[list[str], int]:
    """関連語の前後 width 文字を抜粋する(重なる範囲は結合)。(抜粋, 全ヒット数) を返す。"""
    if not keywords:
        return [], 0
    hits = sorted(m.span() for pat in _keyword_patterns(keywords) for m in pat.finditer(text))
    spans: list[list[int]] = []
    for a, b in hits:
        a, b = max(0, a - width), min(len(text), b + width)
        if spans and a <= spans[-1][1]:
            spans[-1][1] = max(spans[-1][1], b)
        else:
            spans.append([a, b])
    return [text[a:b].replace("\n", " ") for a, b in spans[:max_snippets]], len(hits)


# 作業: 請求項も明細書も無く、要約だけの特許か(Web から取れず、CSV の要約で代用した場合など)を判定する。
def is_abstract_only(rec: dict) -> bool:
    return not rec.get("claims_nfkc") and not rec.get("description_nfkc")


# 作業: 調べたいこと・判定基準・特許本文を、1 つのプロンプトに組み立てる。
def build_prompt(topic: Topic, doc_no: str, rec: dict, max_chars: int = 6000) -> tuple[str, int]:
    """rec は全文JSON(*_nfkc キーを使う)。max_chars で本文の長さを抑える。(プロンプト, 関連語ヒット数)"""
    # 要約しか無い場合(Web から本文が取れなかった場合)は、要約から関連語を探す
    abstract_only = is_abstract_only(rec)
    desc = rec["description_nfkc"] or rec["abstract_nfkc"]
    snippets, hits = keyword_snippets(desc, topic.keywords)
    body = "\n---\n".join(snippets) if snippets else desc[:1500]
    claims = rec["claims_nfkc"]
    prompt = USER_TEMPLATE.format(
        name=topic.name, definition=topic.definition.strip(),
        keywords="、".join(topic.keywords) or "(なし)", include=topic.include.strip() or "(指定なし)",
        exclude=topic.exclude.strip() or "(指定なし)", doc_no=doc_no, title=rec["title_nfkc"],
        abstract=rec["abstract_nfkc"][:1500],
        claims=claims[: max_chars // 2] if claims else "(取得できなかったため、なし)",
        description=("(取得できなかったため、なし。要約だけで判定する)" if abstract_only else body[: max_chars // 2]),
    )
    if abstract_only:
        prompt += "\n\n注意: この特許は要約しか入手できていない。要約に書かれていないことは判断できないので、根拠が弱い場合は score を 0 か 1 にする。"
    return prompt, hits


# ---- バックエンド: (system, user, schema) -> JSON 文字列 -------------------------------

# 作業: Ollama(この PC)に問い合わせる関数を作る。出力は JSON スキーマで制約される。
def ollama_backend(model: str = "qwen3:14b", host: str = "http://localhost:11434", num_ctx: int = 8192) -> Callable:
    import requests

    def run(system: str, user: str, schema: dict) -> str:
        r = requests.post(f"{host}/api/chat", timeout=600, json={
            "model": model, "stream": False, "think": False, "format": schema,
            "options": {"temperature": 0, "num_ctx": num_ctx},
            "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}],
        })
        r.raise_for_status()
        return r.json()["message"]["content"]

    return run


# 作業: guidance(Hugging Face のモデル、サーバー用)に問い合わせる関数を作る。出力は JSON 文法で制約される。
def guidance_backend(lm) -> Callable:
    """lm: guidance のモデル。例: guidance.models.Transformers("Qwen/Qwen3-14B", device_map="auto")。"""
    from guidance import assistant, system as system_role, user as user_role
    from guidance import json as gjson

    def run(system: str, user: str, schema: dict) -> str:
        m = lm
        with system_role():
            m += system
        with user_role():
            m += user
        with assistant():
            m += gjson(name="out", schema=schema)
        return m["out"]

    return run


# ---- 評価 ----------------------------------------------------------------------------

# 作業: LLM の引用が、原文(全角半角・空白の違いを無視)に含まれるかを確認する。
def check_evidence(evidence: list[str], source: str) -> str:
    """引用が原文(NFKC・空白除去)に含まれるかを確認する。"""
    norm = lambda s: re.sub(r"\s+", "", nfkc(s))
    if not evidence:
        return "引用なし"
    src = norm(source)
    ok = sum(norm(e) in src for e in evidence)
    return "全て原文に存在" if ok == len(evidence) else f"原文に無い引用あり({len(evidence) - ok}/{len(evidence)})"


# 作業: 1 件を評価する(プロンプト作成 → LLM → JSON の読み取り → 引用の検証)。
def evaluate_one(backend: Callable, topic: Topic, doc_no: str, rec: dict, max_chars: int = 6000) -> dict:
    prompt, hits = build_prompt(topic, doc_no, rec, max_chars)
    try:
        out = json.loads(backend(SYSTEM, prompt, SCHEMA))
    except (json.JSONDecodeError, KeyError, ValueError) as e:
        return {"score": None, "judgement": "判定失敗", "reason": f"{type(e).__name__}: {e}", "summary": "",
                "evidence": [], "evidence_check": "", "keyword_hits": hits}
    source = rec["abstract_nfkc"] + rec["claims_nfkc"] + rec["description_nfkc"]
    out["evidence_check"] = check_evidence(out.get("evidence", []), source)
    out["keyword_hits"] = hits
    if is_abstract_only(rec):  # 結果の表で、要約だけで判定したと分かるようにする
        out["reason"] = "【要約のみで判定】" + out.get("reason", "")
    return out


# 最終の Excel に出す列と順番。
FINAL_COLUMNS = ["文献番号", "判定", "関連度", "関連語ヒット数", "理由", "引用の検証", "根拠の引用",
                 "発明の名称", "出願人・権利者", "出願企業", "ステータス", "発明内容概要"]

# 法人格・会社の種類を表す語。名前の末尾から取り除く(途中の「インク」等は残すため、末尾だけ)。
_CORP_SUFFIXES = [
    "インコーポレイテッド", "インコーポレイティド", "インコーポレーテッド", "インコーポレイテツド", "コーポレーション", "コーポレイション",
    "ライアビリティ", "リミテッド", "リミティド", "カンパニー", "エルエルシー", "ゲーエムベーハー", "アクチェンゲゼルシャフト",
    "ソシエテアノニム", "アーゲー", "インク",
]
# 法人格を表す漢字の語。名前のどこにあっても取り除く(前株・後株の両方に対応)。
_CORP_KANJI = ["株式会社", "有限会社", "合同会社", "股份有限公司", "有限公司"]


def company_name(applicant, aliases: dict | None = None) -> str:
    """出願人の欄から、出願企業名を作る。

    全角半角をそろえ、区切り(空白・中黒・読点)と法人格の語(株式会社、インコーポレイテッド など)を取り除く。
    aliases({正規化した名前: 表示名})に一致すれば、その表示名を使う。
    出願人が複数の場合、J-PlatPat の一覧に載っているのは先頭の 1 者だけ(「他あり」)。
    """
    if not isinstance(applicant, str) or not applicant.strip():
        return ""
    # 特許庁の外字表記「▲ふん▼」は「份」(股份有限公司)のこと
    name = nfkc(applicant).replace("▲ふん▼", "份")
    for w in _CORP_KANJI:
        name = name.replace(w, "")
    name = re.sub(r"[\s・･,，.．、]+", "", name)
    # 末尾の法人格の語を、なくなるまで繰り返し取り除く(名前全体が消える場合は止める)
    changed = True
    while changed:
        changed = False
        for w in _CORP_SUFFIXES:
            if name.endswith(w) and len(name) > len(w):
                name = name[: -len(w)]
                changed = True
    return (aliases or {}).get(name, name)


# 作業: 文献番号表の各行を、全文 JSON(text_dir/{ID}.json)で評価し、FINAL_COLUMNS の順の結果表を返す。
def evaluate_table(backend: Callable, topic: Topic, numbers, text_dir, max_chars: int = 6000, aliases: dict | None = None):
    from pathlib import Path

    import pandas as pd

    rows = []
    for _, r in numbers.iterrows():
        pid = r.get("Google Patents ID(推定)")
        path = Path(text_dir) / f"{pid}.json"
        applicant = r.get("出願人/権利者")
        more = "（他あり）" if r.get("出願人(他あり)") is True or str(r.get("出願人(他あり)")) == "True" else ""
        base = {"文献番号": r["文献番号"], "リンク": r.get("Google Patents URL(推定)"), "発明の名称": r.get("発明の名称"),
                "出願人・権利者": f"{applicant}{more}" if isinstance(applicant, str) else "",
                "出願企業": company_name(applicant, aliases), "ステータス": r.get("ステータス")}
        if not path.exists():
            rows.append({**base, "判定": "本文なし(未取得)"})
            continue
        rec = json.loads(path.read_text(encoding="utf-8"))
        # 要約で代用した行は Google Patents のページが無い(404 など)ので、リンクを J-PlatPat の URL に替える
        jp_url = r.get("J-PlatPat URL")
        if rec.get("source") == "csv_abstract" and isinstance(jp_url, str) and jp_url:
            base["リンク"] = jp_url
        try:
            o = evaluate_one(backend, topic, r["文献番号"], rec, max_chars)
        # 接続断などでも、1 件の失敗で全体を止めず、失敗として記録して次へ進む
        except Exception as e:
            o = {"score": None, "judgement": "判定失敗", "reason": f"{type(e).__name__}: {e}", "summary": "",
                 "evidence": [], "evidence_check": "", "keyword_hits": None}
        # LLM の概要が空なら、Google Patents の要約(公式)の先頭で代用し、その旨を付ける
        summary = o.get("summary") or (("（要約より）" + rec["abstract_nfkc"][:150]) if rec.get("abstract_nfkc") else "")
        rows.append({**base, "関連度": o["score"], "判定": o["judgement"], "理由": o["reason"],
                     "根拠の引用": " / ".join(o["evidence"]), "引用の検証": o["evidence_check"],
                     "関連語ヒット数": o["keyword_hits"], "発明内容概要": summary})
        print(r["文献番号"], o["score"], o["judgement"], flush=True)
    df = pd.DataFrame(rows)
    for c in FINAL_COLUMNS + ["リンク"]:
        if c not in df:
            df[c] = None
    df = df.sort_values("関連度", ascending=False, na_position="last", kind="stable")
    return df[FINAL_COLUMNS + ["リンク"]].reset_index(drop=True)


# 作業: 結果表を Excel に保存する。文献番号をクリックでページが開くリンクにし(Google Patents。要約で代用した行は J-PlatPat)、見やすく整える。
def save_evaluation_excel(df, path) -> None:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    # LLM の出力に制御文字が混ざると Excel に書けず、ファイルが壊れるので、先に取り除く
    out = df[FINAL_COLUMNS].copy()
    for c in out.columns:
        out[c] = out[c].map(lambda v: ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v)
    with __import__("pandas").ExcelWriter(path, engine="openpyxl") as w:
        out.to_excel(w, index=False, sheet_name="評価結果")
        ws = w.sheets["評価結果"]
        widths = {"文献番号": 18, "判定": 14, "関連度": 8, "関連語ヒット数": 10, "理由": 50, "引用の検証": 20,
                  "根拠の引用": 60, "発明の名称": 40, "出願人・権利者": 28, "出願企業": 20, "ステータス": 22, "発明内容概要": 60}
        wrap = {"理由", "根拠の引用", "発明の名称", "出願人・権利者", "発明内容概要"}
        for j, c in enumerate(FINAL_COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(j)].width = widths[c]
            head = ws.cell(row=1, column=j)
            head.font = Font(bold=True)
            head.fill = PatternFill("solid", fgColor="DDDDDD")
            for i in range(2, len(out) + 2):
                ws.cell(row=i, column=j).alignment = Alignment(wrap_text=c in wrap, vertical="top")
        # 文献番号のセルに、Google Patents へのリンクを付ける
        for i, url in enumerate(df["リンク"], start=2):
            if isinstance(url, str) and url:
                cell = ws.cell(row=i, column=1)
                cell.hyperlink = url
                cell.font = Font(color="0563C1", underline="single")
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
