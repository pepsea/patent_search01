import json
import os

import pandas as pd

from patent_search import fetch_google
from patent_search.evaluate import SCHEMA, Topic, build_prompt, evaluate_one, is_abstract_only
from patent_search.platpat import _clean_abstract, _split_fi, build_list, parse_csv_text

# J-PlatPat の CSV 出力(実物の形)。要約は複数行、FI はカンマ区切りで「G01N37/00,102」のように中にもカンマがある。
HEAD = '"文献番号","出願番号","出願日","公知日","発明の名称","出願人/権利者","FI","要約","公開番号","公告番号","登録番号","審判番号","その他","ステージ","イベント詳細","文献URL"\n'
ROW1 = ('"特開2011-204261","特願2011-123693","2011/06/01","2011/10/13","リシークエンシング","ザ　ネイビー",'
        '"C12N15/00@A,G01N37/00,102,G06F19/00,622","(57)【要約】      （修正有）\n【課題】課題の文です。\n【解決手段】手段の文です。\n【選択図】なし\n",'
        '"特開2011-204261","","特許5517996","","","特許 消滅","年金不納による特許権の消滅","https://www.j-platpat.inpit.go.jp/c1801/PU/JP-2011-204261/11/ja"\n')
ROW2 = ('"特表2010-510230","特願2009-537271","2007/11/06","2008/05/29","吸入剤","テイコク",'
        '"A61K9/12","(57)【要約】\n本発明は吸入剤を提供する。\n","特表2010-510230","","","","","（出願の）却下・拒絶","出願の拒絶・却下","https://example/2"\n')


def test_split_fi_keeps_commas_inside_fi():
    assert _split_fi("C12N15/00@A,G01N37/00,102,G06F19/00,622,C12Q1/68@A") == [
        "C12N15/00@A", "G01N37/00,102", "G06F19/00,622", "C12Q1/68@A"]


def test_clean_abstract():
    raw = "(57)【要約】      （修正有）\n【課題】Ａ。\n【解決手段】Ｂ。\n【選択図】なし\n"
    assert _clean_abstract(raw) == "【課題】Ａ。\n【解決手段】Ｂ。"
    assert _clean_abstract("(57)【要約】\n本発明は〜\n") == "本発明は〜"


def test_parse_csv():
    rows = parse_csv_text(HEAD + ROW1 + ROW2)
    assert [r.doc_no for r in rows] == ["特開2011-204261", "特表2010-510230"]
    r = rows[0]
    assert r.status == ["特許 消滅", "年金不納による特許権の消滅"] and r.fi[1] == "G01N37/00,102"
    assert r.abstract.startswith("【課題】課題の文です。") and "選択図" not in r.abstract
    assert r.jplatpat_url.endswith("/11/ja") and not r.warnings
    d = r.to_dict()
    assert d["google_patent_id"] == "JP2011204261A"


def test_csv_missing_columns_is_clear_error():
    import pytest

    with pytest.raises(ValueError, match="列が足りません"):
        parse_csv_text('"a","b"\n"1","2"\n')


def test_build_list_mixes_txt_and_csv_and_keeps_abstract(tmp_path):
    txt = tmp_path / "a.txt"
    csv_ = tmp_path / "b.csv"
    # txt の方が新しい(ステータスが新しい)が要約は無い → 要約は CSV のものを引き継ぐ
    csv_.write_text(HEAD + ROW1, encoding="utf-8-sig")
    txt.write_text("検索結果一覧(国内文献)\nFI\n1\n\n特開2011-204261\n\n特願2011-123693\n\n2011/06/01\n\n2011/10/13\n\n"
                   "リシークエンシング\n\nザ　ネイビー\n\n特許 有効\n年金納付\n\nC12N15/00@A\n", encoding="utf-8")
    os.utime(csv_, (1000, 1000))
    os.utime(txt, (2000, 2000))
    df, rep = build_list(tmp_path, ["*.txt", "*.csv"])
    assert len(df) == 1 and rep["duplicates_removed"] == 1
    row = df.iloc[0]
    assert row["ステータス"] == "特許 有効 / 年金納付" and row["要約(CSV)"].startswith("【課題】")
    assert row["出典ファイル"] == "b.csv ; a.txt" and set(rep["files"]["形式"]) == {"csv", "txt"}
    # 形式を 1 つに絞れる
    only_csv, _ = build_list(tmp_path, "*.csv")
    assert only_csv.iloc[0]["ステータス"].startswith("特許 消滅")


def test_run_falls_back_to_csv_abstract_on_404(tmp_path, monkeypatch):
    monkeypatch.setattr(fetch_google, "fetch", lambda *a, **k: ("", "JP2011204261A", "404"))
    nums = pd.DataFrame([{"文献番号": "特開2011-204261", "Google Patents ID(推定)": "JP2011204261A",
                          "発明の名称": "リシークエンシング", "要約(CSV)": "【課題】課題の文です。"}])
    res = fetch_google.run(nums, tmp_path / "html", 0, tmp_path / "text")
    r = res.iloc[0]
    assert r["取得状況"] == "404" and r["本文の出所"] == "CSVの要約のみ" and "CSVの要約のみ使用" in r["抽出状況"]
    saved = json.loads((tmp_path / "text" / "JP2011204261A.json").read_text(encoding="utf-8"))
    assert saved["source"] == "csv_abstract" and saved["abstract"] == "【課題】課題の文です。" and saved["claims"] == ""


def res_cols(row):
    return [k for k, v in row.items() if not pd.isna(v)]


def test_run_without_csv_abstract_keeps_failure(tmp_path, monkeypatch):
    monkeypatch.setattr(fetch_google, "fetch", lambda *a, **k: ("", "JP1A", "404"))
    nums = pd.DataFrame([{"文献番号": "特開1", "Google Patents ID(推定)": "JP1A", "発明の名称": "x"}])
    r = fetch_google.run(nums, tmp_path / "html", 0, tmp_path / "text").iloc[0]
    # CSV の要約が無ければ代用せず、失敗のまま(本文の出所・抽出状況の列は作られない)
    assert r["取得状況"] == "404" and "抽出状況" not in res_cols(r)
    assert not (tmp_path / "text" / "JP1A.json").exists()


def test_run_fills_missing_web_abstract_from_csv(tmp_path, monkeypatch):
    html = ('<html><body><span itemprop="title">名称</span><section itemprop="claims"><div class="claim" num="1">'
            '<div class="claim-text">請求項</div></div></section><section itemprop="description">'
            '<div class="description-paragraph" num="0001">本文</div></section></body></html>')
    monkeypatch.setattr(fetch_google, "fetch", lambda *a, **k: (html, "JP1A", "200"))
    nums = pd.DataFrame([{"文献番号": "特開1", "Google Patents ID(推定)": "JP1A", "発明の名称": "x", "要約(CSV)": "CSV要約"}])
    r = fetch_google.run(nums, tmp_path / "html", 0, tmp_path / "text").iloc[0]
    assert r["要約"] == "CSV要約" and r["本文の出所"] == "Google Patents" and "要約はCSVで補完" in r["抽出状況"]
    assert not r["抽出状況"].startswith("欠落")


ABS_ONLY = {"title_nfkc": "リシークエンシング", "abstract_nfkc": "total RNA を解析する方法の要約",
            "claims_nfkc": "", "description_nfkc": ""}
TOPIC = Topic("TOTAL-RNA-seq", "全RNAを対象にしたシーケンス", ["total RNA"])


def test_abstract_only_prompt_and_reason():
    assert is_abstract_only(ABS_ONLY)
    prompt, hits = build_prompt(TOPIC, "特開X", ABS_ONLY)
    assert hits == 1 and "要約しか入手できていない" in prompt and "取得できなかったため" in prompt
    fake = lambda s, u, schema: json.dumps({"score": 1, "judgement": "わずかに関連", "reason": "r", "summary": "s",
                                             "evidence": ["total RNA を解析する方法"]})
    out = evaluate_one(fake, TOPIC, "特開X", ABS_ONLY)
    assert out["reason"] == "【要約のみで判定】r" and out["evidence_check"] == "全て原文に存在"


def test_link_switches_to_jplatpat_for_abstract_substituted_rows(tmp_path):
    from openpyxl import load_workbook

    from patent_search.evaluate import evaluate_table, save_evaluation_excel

    (tmp_path / "JP1A.json").write_text(json.dumps({**ABS_ONLY, "source": "csv_abstract"}), encoding="utf-8")
    web = {"title_nfkc": "t", "abstract_nfkc": "a", "claims_nfkc": "c", "description_nfkc": "d", "source": "web"}
    (tmp_path / "JP2A.json").write_text(json.dumps(web), encoding="utf-8")
    g = "https://patents.google.com/patent/{}/ja"
    nums = pd.DataFrame([
        {"文献番号": "特開1", "Google Patents ID(推定)": "JP1A", "Google Patents URL(推定)": g.format("JP1A"), "J-PlatPat URL": "https://jplatpat/1"},
        {"文献番号": "特開2", "Google Patents ID(推定)": "JP2A", "Google Patents URL(推定)": g.format("JP2A"), "J-PlatPat URL": "https://jplatpat/2"},
        {"文献番号": "特開3", "Google Patents ID(推定)": "JP3A", "Google Patents URL(推定)": g.format("JP3A"), "J-PlatPat URL": "https://jplatpat/3"},
    ])
    fake = lambda s, u, schema: json.dumps({"score": 1, "judgement": "わずかに関連", "reason": "r", "summary": "s", "evidence": []})
    df = evaluate_table(fake, TOPIC, nums, tmp_path)
    links = dict(zip(df["文献番号"], df["リンク"]))
    assert links["特開1"] == "https://jplatpat/1"           # 要約で代用 → J-PlatPat
    assert links["特開2"] == g.format("JP2A")               # Web の本文あり → Google Patents のまま
    assert links["特開3"] == g.format("JP3A")               # 本文なし(未取得) → そのまま
    save_evaluation_excel(df, tmp_path / "o.xlsx")
    ws = load_workbook(tmp_path / "o.xlsx")["評価結果"]
    assert {ws.cell(row=i, column=1).value: ws.cell(row=i, column=1).hyperlink.target for i in (2, 3, 4)}["特開1"] == "https://jplatpat/1"


def test_fulltext_json_is_saved_under_the_list_id_even_when_fallback_id_worked(tmp_path, monkeypatch):
    # 一覧の ID は JP1A。取得は種別コードなしの JP1 で成功した場合でも、評価が探す JP1A.json で保存される
    html = ('<html><body><span itemprop="title">名称</span><section itemprop="claims"><div class="claim" num="1">'
            '<div class="claim-text">請求項</div></div></section><section itemprop="description">'
            '<div class="description-paragraph" num="0001">本文</div></section></body></html>')
    monkeypatch.setattr(fetch_google, "fetch", lambda *a, **k: (html, "JP1", "200"))
    nums = pd.DataFrame([{"文献番号": "特開1", "Google Patents ID(推定)": "JP1A", "発明の名称": "x"}])
    res = fetch_google.run(nums, tmp_path / "html", 0, tmp_path / "text").iloc[0]
    assert (tmp_path / "text" / "JP1A.json").exists() and not (tmp_path / "text" / "JP1.json").exists()
    assert res["使用ID"] == "JP1" and res["全文ファイル"].endswith("JP1A.json")


def test_check_texts_reports_and_suggests_other_run_folder(tmp_path, capsys):
    import pytest

    from patent_search.evaluate import check_texts

    old_run, new_run = tmp_path / "res" / "T_20261005_100000", tmp_path / "res" / "T_20261005_110000"
    (old_run / "text").mkdir(parents=True)
    (new_run / "text").mkdir(parents=True)
    (old_run / "text" / "JP1A.json").write_text("{}", encoding="utf-8")
    nums = pd.DataFrame([{"文献番号": "特開1", "Google Patents ID(推定)": "JP1A"},
                         {"文献番号": "特開2", "Google Patents ID(推定)": "JP2A"}])
    # 全件に全文 JSON がある → そのまま通る
    assert check_texts(nums.head(1), old_run / "text") == 1
    # 一部だけある → 件数を表示して続行
    capsys.readouterr()  # ここまでの表示を捨てる
    assert check_texts(nums, old_run / "text") == 1
    screen = capsys.readouterr().out
    assert "2 件のうち、全文 JSON がある特許: 1 件" in screen and "特開2" in screen and "評価されず" in screen
    # 1 件もない(セル4の再実行で新しい空の調査フォルダを見ている) → 止まり、全文 JSON がある別フォルダを教える
    with pytest.raises(RuntimeError) as e:
        check_texts(nums, new_run / "text", tmp_path / "res", new_run)
    msg = str(e.value)
    assert "全文 JSON が 1 件もない" in msg and 'RUN_DIR = "T_20261005_100000"' in msg and "T_20261005_110000" not in msg


def test_missing_text_row_has_a_reason(tmp_path, capsys):
    from patent_search.evaluate import evaluate_table

    nums = pd.DataFrame([{"文献番号": "特開9", "Google Patents ID(推定)": "JP9A", "発明の名称": "x"}])
    df = evaluate_table(lambda *a, **k: "{}", Topic("T", "d", ["k"]), nums, tmp_path, stream=True)
    assert df.loc[0, "判定"] == "本文なし(未取得)" and "全文 JSON がありません(JP9A.json)" in df.loc[0, "理由"]
    assert "評価を飛ばしました" in capsys.readouterr().out


def test_failure_reason_is_printed_to_screen(tmp_path, capsys):
    from patent_search.evaluate import evaluate_table

    (tmp_path / "JP1A.json").write_text(json.dumps(ABS_ONLY), encoding="utf-8")
    nums = pd.DataFrame([{"文献番号": "特開1", "Google Patents ID(推定)": "JP1A", "発明の名称": "x"}])

    def boom(*a, **k):
        raise ConnectionError("Ollama に接続できません")

    df = evaluate_table(boom, TOPIC, nums, tmp_path)
    assert df.loc[0, "判定"] == "判定失敗" and "接続できません" in df.loc[0, "理由"]
    assert "判定失敗  理由: ConnectionError: Ollama に接続できません" in capsys.readouterr().out
